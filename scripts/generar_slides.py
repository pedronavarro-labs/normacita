"""Genera las slides (.pptx) a partir de docs/presentacion/GUION-SLIDES.md.

Uso:  pip install python-pptx  &&  python scripts/generar_slides.py
El .pptx se puede importar en Google Slides (subirlo a Drive y abrirlo como Presentaciones).
Diseño sencillo: título, viñetas, capturas de docs/img/ y diagramas con formas editables.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
GUION = ROOT / "docs/presentacion/GUION-SLIDES.md"
SALIDA = ROOT / "docs/presentacion/NormaCita-slides.pptx"
IMG = ROOT / "docs/img"

AZUL = RGBColor(0x1D, 0x4E, 0xD8)
AZUL_OSCURO = RGBColor(0x1E, 0x3A, 0x8A)
AZUL_SUAVE = RGBColor(0xE8, 0xEF, 0xFF)
TEXTO = RGBColor(0x17, 0x20, 0x33)
GRIS = RGBColor(0x4A, 0x55, 0x68)
BORDE = RGBColor(0xD6, 0xDC, 0xE6)
AMBAR = RGBColor(0xF0, 0xB4, 0x29)
AMBAR_SUAVE = RGBColor(0xFF, 0xF7, 0xE6)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Calibri"

W, H = Inches(13.333), Inches(7.5)


def parse_guion(texto: str) -> list[dict]:
    slides = []
    for bloque in re.split(r"^## Slide \d+ · ", texto, flags=re.M)[1:]:
        titulo, _, cuerpo = bloque.partition("\n")
        contenido = cuerpo.split("**Contenido**", 1)[1].split("**Visual sugerido:**", 1)[0]
        visual = cuerpo.split("**Visual sugerido:**", 1)[1].split("**Notas del orador:**", 1)[0]
        notas = cuerpo.split("**Notas del orador:**", 1)[1]
        slides.append(
            {
                "titulo": titulo.strip(),
                "bullets": [
                    b[2:].strip() for b in contenido.strip().splitlines() if b.startswith("- ")
                ],
                "visual": visual.strip(),
                "notas": notas.strip(),
            }
        )
    return slides


def _texto(frame, texto, size, color=TEXTO, bold=False, align=PP_ALIGN.LEFT):
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = texto
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FUENTE
    return p


def caja(slide, x, y, w, h, texto, size=16, color=TEXTO, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    _texto(tb.text_frame, texto, size, color, bold, align)
    return tb


def forma(
    slide, x, y, w, h, texto="", relleno=AZUL_SUAVE, borde=None, color=TEXTO, size=14, bold=False
):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.12
    s.fill.solid()
    s.fill.fore_color.rgb = relleno
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if texto:
        _texto(tf, texto, size, color, bold, PP_ALIGN.CENTER)
    return s


def flecha(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = GRIS
    c.line.width = Pt(2)
    c.line._get_or_add_ln().append(  # punta de flecha
        c.line._get_or_add_ln().makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}
        )
    )


def imagen(slide, nombre, x, y, w=None, h=None, borde=True):
    ruta = IMG / nombre
    if not ruta.exists():
        return None
    pic = slide.shapes.add_picture(str(ruta), x, y, width=w, height=h)
    if borde:
        pic.line.color.rgb = BORDE
        pic.line.width = Pt(1)
    return pic


def base(prs, datos, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    barra.fill.solid()
    barra.fill.fore_color.rgb = AZUL
    barra.line.fill.background()
    caja(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.9), datos["titulo"], 34, TEXTO, True)
    pie = f"NormaCita · Pedro Navarro Arocha · TFM Máster en Desarrollo con IA · {n}/{total}"
    caja(s, Inches(0.6), Inches(6.95), Inches(12), Inches(0.4), pie, 11, GRIS)
    s.notes_slide.notes_text_frame.text = datos["notas"]
    return s


def viñetas(slide, bullets, x, y, w, h, size=20):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = "▸  "
        r1.font.color.rgb, r1.font.size, r1.font.name = AZUL, Pt(size), FUENTE
        r2 = p.add_run()
        r2.text = b.replace("`", "")
        r2.font.color.rgb, r2.font.size, r2.font.name = TEXTO, Pt(size), FUENTE
        if "[PENDIENTE" in b:
            r2.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
            r2.font.italic = True
    return tb


def portada(prs, d, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = AZUL_OSCURO
    fondo.line.fill.background()
    logo = forma(
        s, Inches(0.8), Inches(1.2), Inches(1.0), Inches(1.0), "N", BLANCO, None, AZUL, 40, True
    )
    logo.adjustments[0] = 0.25
    caja(s, Inches(0.8), Inches(2.4), Inches(6.2), Inches(1.2), d["titulo"], 54, BLANCO, True)
    caja(
        s, Inches(0.8), Inches(3.5), Inches(6.2), Inches(0.8), d["bullets"][0], 26, AZUL_SUAVE, True
    )
    caja(s, Inches(0.8), Inches(4.3), Inches(6.0), Inches(1.0), d["bullets"][1], 18, AZUL_SUAVE)
    caja(s, Inches(0.8), Inches(5.9), Inches(6.2), Inches(0.9), d["bullets"][2], 14, AZUL_SUAVE)
    imagen(s, "01-inicio.png", Inches(7.3), Inches(1.0), w=Inches(5.6), borde=False)
    s.notes_slide.notes_text_frame.text = d["notas"]


def slide_problema(s):
    for i, (num, txt) in enumerate(
        [("29", "artículos en el Real Decreto"), ("52", "instrucciones técnicas ITC-BT")]
    ):
        x = Inches(8.4)
        y = Inches(1.6 + i * 2.2)
        forma(s, x, y, Inches(4.2), Inches(1.9), relleno=AZUL_SUAVE)
        caja(s, x, y + Inches(0.15), Inches(4.2), Inches(1.0), num, 54, AZUL, True, PP_ALIGN.CENTER)
        caja(
            s, x, y + Inches(1.15), Inches(4.2), Inches(0.6), txt, 16, GRIS, False, PP_ALIGN.CENTER
        )


def slide_usuarios(s, bullets):
    for i, b in enumerate(bullets):
        quien, _, que = b.partition(":")
        x = Inches(0.6 + i * 4.15)
        forma(s, x, Inches(1.7), Inches(3.9), Inches(4.6), relleno=AZUL_SUAVE)
        caja(
            s,
            x + Inches(0.25),
            Inches(1.95),
            Inches(3.4),
            Inches(1.0),
            quien.strip(),
            22,
            AZUL_OSCURO,
            True,
        )
        caja(
            s,
            x + Inches(0.25),
            Inches(3.1),
            Inches(3.4),
            Inches(3.0),
            que.strip().capitalize(),
            18,
            TEXTO,
        )


def slide_flujo(s):
    pasos = [
        "Pregunta",
        "BM25\n(832 fragmentos)",
        "¿Umbral de\npuntuación y\ncobertura?",
        "LLM o\nmodo demo",
        "Validar\ncitas [n]",
        "Respuesta\ncon citas",
    ]
    w, h, y = Inches(1.85), Inches(1.25), Inches(5.15)
    gap = Inches(0.27)
    x0 = Inches(0.6)
    for i, p in enumerate(pasos):
        x = x0 + i * (w + gap)
        relleno = AMBAR_SUAVE if i == 2 else (AZUL if i == 5 else AZUL_SUAVE)
        color = BLANCO if i == 5 else TEXTO
        forma(s, x, y, w, h, p, relleno, AMBAR if i == 2 else None, color, 14, True)
        if i < len(pasos) - 1:
            flecha(s, x + w, y + h // 2, x + w + gap, y + h // 2)
    xs = x0 + 2 * (w + gap)
    forma(
        s,
        xs,
        Inches(6.55) - Inches(0.05),
        w,
        Inches(0.45),
        "No → «Sin base»",
        AMBAR_SUAVE,
        AMBAR,
        TEXTO,
        12,
        True,
    )


def slide_capas(s):
    x, w = Inches(8.0), Inches(4.8)
    capas = [
        ("api · FastAPI, validación, seguridad, UI", AZUL_SUAVE),
        ("application · AskQuestion + puertos", AZUL_SUAVE),
        ("domain · Fragment, Citation, Answer", AZUL),
        ("infrastructure · BM25, corpus, LLM fake / OpenAI-compatible", AZUL_SUAVE),
    ]
    for i, (t, c) in enumerate(capas):
        forma(
            s,
            x,
            Inches(1.6 + i * 1.2),
            w,
            Inches(0.95),
            t,
            c,
            None,
            BLANCO if c == AZUL else TEXTO,
            14,
            True,
        )


def slide_metricas(s):
    filas = [
        ("Métrica (52 preguntas)", "Antes", "Después"),
        ("hit@1", "0,548", "0,714"),
        ("hit@3", "0,738", "0,952"),
        ("Rechazo correcto", "0,30", "1,00"),
        ("Validación hit@1 / hit@3", "0,333 / 0,667", "0,50 / 0,917"),
    ]
    t = s.shapes.add_table(len(filas), 3, Inches(8.0), Inches(1.7), Inches(4.9), Inches(2.6)).table
    t.columns[0].width = Inches(2.5)
    t.columns[1].width = Inches(1.2)
    t.columns[2].width = Inches(1.2)
    for r, fila in enumerate(filas):
        for c, val in enumerate(fila):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = AZUL if r == 0 else (AZUL_SUAVE if r % 2 else BLANCO)
            _texto(
                cell.text_frame,
                val,
                13,
                BLANCO if r == 0 else TEXTO,
                r == 0 or c == 2,
                PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER,
            )
    caja(
        s,
        Inches(8.0),
        Inches(4.5),
        Inches(4.9),
        Inches(1.2),
        "Fuente: docs/EVALUACION.md · mismo corpus y mismas preguntas antes y después.",
        12,
        GRIS,
    )


def slide_ci(s):
    pasos = ["git push", "ruff", "pytest +\nevaluación", "Docker +\n/health", "Render"]
    w, h, y, gap = Inches(2.15), Inches(1.0), Inches(5.3), Inches(0.38)
    for i, p in enumerate(pasos):
        x = Inches(0.6) + i * (w + gap)
        forma(
            s,
            x,
            y,
            w,
            h,
            p,
            AZUL if i == 4 else AZUL_SUAVE,
            None,
            BLANCO if i == 4 else TEXTO,
            15,
            True,
        )
        if i < len(pasos) - 1:
            flecha(s, x + w, y + h // 2, x + w + gap, y + h // 2)


def main() -> None:
    slides = parse_guion(GUION.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    total = len(slides)
    ancho_txt = Inches(7.1)
    for n, d in enumerate(slides, start=1):
        if n == 1:
            portada(prs, d, total)
            continue
        s = base(prs, d, n, total)
        x, y = Inches(0.6), Inches(1.45)
        if n == 2:
            viñetas(s, d["bullets"], x, y, ancho_txt, Inches(5.2))
            slide_problema(s)
        elif n == 3:
            slide_usuarios(s, d["bullets"])
        elif n == 4:
            viñetas(s, d["bullets"], x, y, Inches(5.6), Inches(3.2), 18)
            imagen(s, "02b-respuesta-detalle.png", Inches(6.4), Inches(1.3), h=Inches(5.45))
            imagen(s, "03b-sin-base-detalle.png", Inches(0.6), Inches(4.75), w=Inches(5.4))
        elif n == 5:
            viñetas(s, d["bullets"], x, y, Inches(12.2), Inches(3.6), 17)
            slide_flujo(s)
        elif n == 6:
            viñetas(s, d["bullets"], x, y, ancho_txt, Inches(5.2), 18)
            slide_capas(s)
        elif n == 8:
            viñetas(s, d["bullets"], x, y, ancho_txt, Inches(5.2), 18)
            slide_metricas(s)
        elif n == 9:
            viñetas(s, d["bullets"], x, y, Inches(12.2), Inches(3.6), 19)
            slide_ci(s)
        elif n == 12:
            viñetas(s, d["bullets"], x, y, Inches(8.0), Inches(5.2), 20)
            imagen(s, "04a-movil-inicio.png", Inches(9.6), Inches(1.3), h=Inches(5.5))
        else:
            viñetas(s, d["bullets"], x, y, Inches(12.0), Inches(5.3), 20)
    prs.save(SALIDA)
    print(f"{SALIDA.relative_to(ROOT)}: {total} slides")


if __name__ == "__main__":
    main()
