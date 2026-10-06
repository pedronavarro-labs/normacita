"""Genera la memoria en .docx (y PDF opcional) a partir de docs/MEMORIA.md.

- Renderiza cada bloque Mermaid precedido de `<!-- diagrama: nombre | título -->` a
  docs/img/diagramas/<nombre>.png con mermaid-cli (`mmdc`).
- Convierte con pandoc a .docx (importable en Google Docs) y, con --pdf, a PDF
  vía HTML + Chrome headless (Playwright).

Requisitos (solo para regenerar; no son dependencias de la app):
  pandoc, python-docx, mermaid-cli (npm i @mermaid-js/mermaid-cli) y, para PDF, playwright.
Uso: python scripts/generar_memoria.py --salida /ruta/carpeta [--pdf] [--mmdc /ruta/mmdc]
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
MEMORIA = DOCS / "MEMORIA.md"
DIAGRAMAS = DOCS / "img" / "diagramas"
BLOQUE = re.compile(r"<!-- diagrama: ([\w-]+) \| ([^>]+?) -->\n```mermaid\n(.*?)```", re.S)
ANCHOS = {"04a-movil-inicio.png": "38%", "03b-sin-base-detalle.png": "75%"}

CSS = """
@page { size: A4; margin: 2cm 1.8cm; }
body { font-family: "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif; font-size: 10.5pt;
       line-height: 1.5; color: #172033; max-width: none; }
h1 { color: #1e3a8a; font-size: 22pt; border-bottom: 3px solid #1d4ed8; padding-bottom: 6px; }
h2 { color: #1d4ed8; font-size: 15pt; margin-top: 1.6em; page-break-after: avoid; }
h2#anexos, h2#resumen { page-break-before: always; }
h3 { color: #1e3a8a; font-size: 12pt; page-break-after: avoid; }
table { border-collapse: collapse; width: 100%; margin: .8em 0; font-size: 9.5pt;
        page-break-inside: avoid; }
th { background: #1d4ed8; color: #fff; text-align: left; }
th, td { border: 1px solid #d6dce6; padding: 5px 7px; vertical-align: top; }
tr:nth-child(even) td { background: #f5f7fb; }
blockquote { background: #fff7e6; border-left: 4px solid #f0b429; margin: 1em 0;
             padding: .5em 1em; }
code { background: #eef1f6; padding: 0 3px; border-radius: 3px; font-size: 9pt; }
pre { background: #f5f7fb; border: 1px solid #d6dce6; padding: 8px; font-size: 8.5pt;
      white-space: pre-wrap; }
figure { text-align: center; margin: 1em 0; page-break-inside: avoid; }
figure img { max-width: 100%; border: 1px solid #d6dce6; }
figcaption { color: #4a5568; font-size: 9pt; font-style: italic; }
"""


def render_diagramas(md: str, mmdc: list[str]) -> str:
    DIAGRAMAS.mkdir(parents=True, exist_ok=True)
    cfg = Path(tempfile.mkdtemp()) / "puppeteer.json"
    chrome = shutil.which("google-chrome") or shutil.which("chromium")
    cfg.write_text(f'{{"executablePath": "{chrome}", "args": ["--no-sandbox"]}}', encoding="utf-8")

    def sustituir(m: re.Match[str]) -> str:
        nombre, titulo, codigo = m.group(1), m.group(2).strip(), m.group(3)
        src = cfg.parent / f"{nombre}.mmd"
        src.write_text(codigo, encoding="utf-8")
        png = DIAGRAMAS / f"{nombre}.png"
        cmd = [*mmdc, "-i", str(src), "-o", str(png), "-s", "2", "-b", "white", "-p", str(cfg)]
        subprocess.run(cmd, check=True, capture_output=True)  # noqa: S603
        return f"![{titulo}](img/diagramas/{nombre}.png){{width=90%}}"

    return BLOQUE.sub(sustituir, md)


def preparar(md: str) -> str:
    def ancho(m: re.Match[str]) -> str:
        alt, ruta = m.group(1), m.group(2)
        w = ANCHOS.get(Path(ruta).name, "85%")
        return f"![{alt}]({ruta}){{width={w}}}"

    return re.sub(r"!\[([^\]]*)\]\((img/[^)]+\.png)\)(?!\{)", ancho, md)


def estilizar_docx(ruta: Path) -> None:
    """Bordes y cabecera azul en las tablas, para que se vean bien en Google Docs."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document(str(ruta))
    estilos = {s.name.lower(): s for s in doc.styles if s.name}
    for nombre in ("normal", "body text", "first paragraph", "compact"):
        if nombre in estilos:
            estilos[nombre].font.name = "Calibri"
            estilos[nombre].font.size = Pt(11)
    for nivel, color in ((1, "1E3A8A"), (2, "1D4ED8"), (3, "1E3A8A")):
        if (st := estilos.get(f"heading {nivel}")) is not None:
            st.font.name = "Calibri"
            st.font.color.rgb = RGBColor.from_string(color)
    for tabla in doc.tables:
        tbl_pr = tabla._tbl.tblPr
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "D6DCE6")
            bordes.append(b)
        tbl_pr.append(bordes)
        for celda in tabla.rows[0].cells:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "E8EFFF")
            celda._tc.get_or_add_tcPr().append(shd)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.bold = True
    doc.save(str(ruta))


def a_pdf(md_tmp: Path, salida: Path) -> None:
    html = salida.with_suffix(".html")
    css = md_tmp.parent / "memoria.css"
    css.write_text(CSS, encoding="utf-8")
    cmd = ["pandoc", str(md_tmp), "-f", "markdown+gfm_auto_identifiers", "-s", "--embed-resources",
           "--resource-path", str(DOCS), "-c", str(css), "--metadata", "title=NormaCita · Memoria",
           "-o", str(html)]  # fmt: skip
    subprocess.run(cmd, check=True)  # noqa: S603
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome")
        pg = b.new_page()
        pg.goto(html.resolve().as_uri())
        pg.add_style_tag(content="header#title-block-header { display: none; }")
        pie = (
            '<div style="font-size:8px;width:100%;text-align:center;color:#4a5568">'
            "NormaCita · Memoria del TFM · "
            '<span class="pageNumber"></span>/<span class="totalPages"></span></div>'
        )
        margen = {lado: "1.6cm" for lado in ("top", "bottom", "left", "right")}
        pg.pdf(
            path=str(salida),
            format="A4",
            print_background=True,
            display_header_footer=True,
            header_template="<span></span>",
            footer_template=pie,
            margin=margen,
        )
        b.close()
    html.unlink()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--salida", type=Path, required=True, help="Carpeta de salida (fuera del repo)")
    ap.add_argument("--pdf", action="store_true", help="Generar también el PDF")
    ap.add_argument("--mmdc", default="mmdc", help="Ruta al ejecutable de mermaid-cli")
    args = ap.parse_args()
    args.salida.mkdir(parents=True, exist_ok=True)

    md = render_diagramas(MEMORIA.read_text(encoding="utf-8"), [args.mmdc])
    md = preparar(md)
    tmp = Path(tempfile.mkdtemp()) / "MEMORIA.md"
    tmp.write_text(md, encoding="utf-8")

    docx = args.salida / "NormaCita-MEMORIA.docx"
    cmd = ["pandoc", str(tmp), "-f", "markdown+gfm_auto_identifiers", "--resource-path", str(DOCS),
           "-o", str(docx)]  # fmt: skip
    subprocess.run(cmd, check=True)  # noqa: S603
    estilizar_docx(docx)
    print(docx)
    if args.pdf:
        pdf = args.salida / "NormaCita-MEMORIA.pdf"
        a_pdf(tmp, pdf)
        print(pdf)


if __name__ == "__main__":
    main()
